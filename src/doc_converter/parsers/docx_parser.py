"""
DOCX (Microsoft Word) document parser.
"""

import os
from typing import Dict, Any
from .base_parser import BaseParser


class DocxParser(BaseParser):
    """Parser for DOCX (Microsoft Word) documents."""
    
    def __init__(self):
        super().__init__()
        self.supported_formats = ['docx']
        self._python_docx = None
    
    def _import_docx(self):
        """Lazy import of python-docx library."""
        if self._python_docx is None:
            try:
                import docx
                self._python_docx = docx
            except ImportError:
                raise ImportError(
                    "python-docx is required for DOCX parsing. "
                    "Install it with: pip install python-docx"
                )
        return self._python_docx
    
    def can_parse(self, file_path: str) -> bool:
        """Check if file is a DOCX document."""
        return file_path.lower().endswith('.docx')
    
    def parse(self, file_path: str) -> Dict[str, Any]:
        """
        Parse DOCX document.
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            Standardized document structure
        """
        docx_module = self._import_docx()
        doc = docx_module.Document(file_path)
        
        result = self._create_base_structure('docx')
        
        # Extract metadata
        self._extract_metadata(doc, result['metadata'], file_path)
        
        # Parse content
        position = 0
        for element in doc.element.body:
            tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag
            
            if tag == 'p':
                # Process paragraph or heading
                para = None
                for p in doc.paragraphs:
                    if p._element == element:
                        para = p
                        break
                
                if para:
                    if para.style.name.startswith('Heading'):
                        self._parse_heading(para, result['content'], position)
                    else:
                        self._parse_paragraph(para, result['content'], position)
                    position += 1
            
            elif tag == 'tbl':
                # Process table
                table = None
                for t in doc.tables:
                    if t._element == element:
                        table = t
                        break
                
                if table:
                    self._parse_table(table, result['content'], position)
                    position += 1
        
        # Handle images
        self._parse_images(doc, result['content'])
        
        return result
    
    def _extract_metadata(self, doc, metadata: Dict[str, Any], 
                         file_path: str) -> None:
        """Extract document metadata."""
        core_props = doc.core_properties
        
        if core_props.title:
            metadata['title'] = core_props.title
        if core_props.author:
            metadata['author'] = core_props.author
        if core_props.created:
            metadata['created_date'] = core_props.created.isoformat()
        if core_props.modified:
            metadata['modified_date'] = core_props.modified.isoformat()
        
        # Word count
        word_count = sum(len(para.text.split()) for para in doc.paragraphs)
        metadata['word_count'] = word_count
    
    def _parse_heading(self, para, content: list, position: int) -> None:
        """Parse heading element."""
        # Extract heading level from style name
        style_name = para.style.name
        level = 1
        if 'Heading' in style_name:
            try:
                level = int(style_name.split()[-1])
            except (ValueError, IndexError):
                level = 1
        
        style_info = self._extract_style(para)
        
        self._add_content_item(
            content, 'heading', position,
            level=level,
            text=para.text,
            style=style_info
        )
    
    def _parse_paragraph(self, para, content: list, position: int) -> None:
        """Parse paragraph element."""
        if not para.text.strip():
            return
        
        style_info = self._extract_style(para)
        
        self._add_content_item(
            content, 'paragraph', position,
            text=para.text,
            style=style_info
        )
    
    def _parse_table(self, table, content: list, position: int) -> None:
        """Parse table element."""
        rows_data = []
        headers = []
        
        for i, row in enumerate(table.rows):
            row_data = [cell.text for cell in row.cells]
            if i == 0:
                headers = row_data
            rows_data.append(row_data)
        
        self._add_content_item(
            content, 'table', position,
            headers=headers,
            rows=rows_data
        )
    
    def _parse_images(self, doc, content: list) -> None:
        """Parse image elements."""
        position = len(content)
        
        # Images are in relationships
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                image_part = rel.target_part
                
                self._add_content_item(
                    content, 'image', position,
                    src=os.path.basename(rel.target_ref),
                    alt=f"Image {position}"
                )
                position += 1
    
    def _extract_style(self, para) -> Dict[str, Any]:
        """Extract style information from paragraph."""
        style = {}
        
        if para.runs:
            first_run = para.runs[0]
            if first_run.bold:
                style['bold'] = True
            if first_run.italic:
                style['italic'] = True
            if first_run.underline:
                style['underline'] = True
            if first_run.font.size:
                style['font_size'] = first_run.font.size.pt
            if first_run.font.name:
                style['font_family'] = first_run.font.name
        
        return style
