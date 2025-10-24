from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import ns
from docx.oxml.shared import OxmlElement


@dataclass
class SampleDocumentDetails:
    path: Path
    image_path: Path


def create_sample_docx(destination: Path) -> SampleDocumentDetails:
    document = Document()
    document.core_properties.title = "Sample Document"
    document.core_properties.author = "Author A"
    document.core_properties.keywords = "alpha, beta"

    document.add_heading("Main Heading", level=1)
    paragraph = document.add_paragraph("This is a paragraph with ")
    bold_run = paragraph.add_run("bold text")
    bold_run.bold = True
    paragraph.add_run(" and ")
    italic_run = paragraph.add_run("italic text")
    italic_run.italic = True
    paragraph.add_run(" plus a ")

    hyperlink_paragraph = document.add_paragraph()
    add_hyperlink(hyperlink_paragraph, "https://example.com", "Example Link")

    document.add_paragraph("Bullet item 1", style="List Bullet")
    document.add_paragraph("Bullet item 2", style="List Bullet")
    document.add_paragraph("Number item 1", style="List Number")
    document.add_paragraph("Number item 2", style="List Number")

    table = document.add_table(rows=2, cols=2)
    table.style = "Light List"
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)

    code = document.add_paragraph()
    code_run = code.add_run("print('hello world')")
    code_run.font.name = "Courier New"

    image_path = destination.parent / "sample-image.png"
    generate_image(image_path)
    document.add_picture(str(image_path))

    document.save(destination)
    return SampleDocumentDetails(path=destination, image_path=image_path)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(ns.qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(ns.qn("w:val"), "single")
    color = OxmlElement("w:color")
    color.set(ns.qn("w:val"), "0000FF")
    r_pr.append(underline)
    r_pr.append(color)
    new_run.append(r_pr)

    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def generate_image(path: Path) -> None:
    image = Image.new("RGB", (200, 120), color=(255, 255, 255))
    draw = ImageDraw.Draw(image)
    draw.rectangle([10, 10, 190, 110], outline="black", width=3)
    draw.text((20, 40), "Sample Image", fill="black")
    image.save(path)


@contextmanager
def libreoffice_conversion(source: Path, target_suffix: str) -> Iterator[Path]:
    from shutil import which
    import subprocess
    import tempfile

    soffice = which("libreoffice") or which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice not found")

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                target_suffix.lstrip("."),
                str(source),
                "--outdir",
                str(tmp),
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        converted = tmp / (source.stem + target_suffix)
        if not converted.exists():
            raise FileNotFoundError(converted)
        yield converted


